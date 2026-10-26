"""In-memory annotated exports. Original document layout is not reconstructed."""

from __future__ import annotations

import csv
import html
import io
import re
from pathlib import Path

import fitz
from docx import Document
from docx.oxml import OxmlElement
from docx.oxml.ns import qn
from docx.shared import Inches, Pt

COLORS = {"C1": "FFF0C2", "C2": "DCECE8"}
FORMAT_NOTICE = "导出为重新排版的标注文本，保留段落与换行；不复原上传文档的原始字体、图片或页面布局。"


def _matches(analysis: dict) -> list[dict]:
    """Use only valid, non-overlapping ranges, even for hand-crafted payloads."""
    text = analysis["text"]
    items = {item["id"]: item for item in analysis.get("items", [])}
    valid = []
    previous_end = 0
    occurrences = sorted(analysis.get("occurrences", []), key=lambda item: item.get("start", -1))
    for occurrence in occurrences:
        start, end = occurrence.get("start"), occurrence.get("end")
        item = items.get(occurrence.get("item_id"))
        if (
            isinstance(start, int) and isinstance(end, int)
            and previous_end <= start < end <= len(text)
            and item and item.get("level") in COLORS
        ):
            valid.append({**occurrence, "item": item})
            previous_end = end
    return valid


def _line_segments(text: str, matches: list[dict]):
    offset = 0
    for line in text.split("\n"):
        segments = []
        cursor = offset
        line_end = offset + len(line)
        for match in matches:
            if match["end"] <= offset:
                continue
            if match["start"] >= line_end:
                break
            start = max(match["start"], offset)
            end = min(match["end"], line_end)
            if start > cursor:
                segments.append((text[cursor:start], None))
            if end > start:
                segments.append((text[start:end], match))
                cursor = end
        if cursor < line_end:
            segments.append((text[cursor:line_end], None))
        yield segments
        offset = line_end + 1


def _basis(item: dict) -> str:
    source = item.get("source") or {}
    if item.get("status") == "reference":
        label = "来源支持；上下文需核对"
    else:
        label = "估计，未经来源验证"
    known_profile = "Octanove" in str(source.get("name", "")) or "openlanguageprofiles/olp-en-cefrj" in str(source.get("url", ""))
    attribution = source.get("attribution") or source.get("creator") or ("Octanove Labs / Open Language Profiles" if known_profile else None)
    license_name = source.get("license") or ("CC BY-SA 4.0" if known_profile else None)
    license_url = source.get("license_url") or (
        "https://creativecommons.org/licenses/by-sa/4.0/" if license_name == "CC BY-SA 4.0" else None
    )
    modifications = source.get("modifications") or (
        "变更：词形匹配、大小写规范化及候选去重；词典定义及用例保留各自来源，中文学习注释与 AI 辅助另行标明。"
        "派生词汇数据按 CC BY-SA 4.0 分享。"
        if known_profile else None
    )
    detail = "\n".join(str(value) for value in (
        source.get("name"), source.get("url"), attribution, license_name, license_url,
        modifications, item.get("notes"),
    ) if value)
    wordset = item.get("example_source") or item.get("definition_source")
    if wordset or item.get("dictionary_senses"):
        wordset = wordset or {}
        detail += "\n" + "\n".join(str(value) for value in (
            "英文定义与用例来源：" + wordset.get("name", "Wordset Dictionary"),
            wordset.get("url", "https://github.com/wordset/wordset-dictionary"),
            wordset.get("creator", "Wordset contributors; derived from Princeton WordNet 3.0"),
            wordset.get("license", "CC BY-SA 4.0"),
            wordset.get("license_url", "https://creativecommons.org/licenses/by-sa/4.0/"),
            "保留 Princeton WordNet 3.0 版权及许可：https://wordnet.princeton.edu/license-and-commercial-use",
            "变更：按词头/词性筛选与组合；义项和 CEFR 对应未经核实。用例包含短语，不保证都是完整句子。",
        ) if value)
    return label + ("\n" + detail if detail else "")


def _definitions(item: dict) -> str:
    if "dictionary_senses" not in item or (not item.get("dictionary_senses") and not item.get("definition_source") and item.get("definition")):
        return str(item.get("definition", ""))
    parts = [f"[{sense.get('pos', '')}] {sense.get('definition', '')}" +
             ("\n同义词参考（需核对语境）：" + " · ".join(sense['synonyms']) if sense.get('synonyms') else '')
             for sense in item.get("dictionary_senses", [])]
    if not parts:
        parts.append("Wordset 暂无匹配义项")
    if item.get("ai_definition"):
        parts.append("AI 辅助定义（待核）：" + str(item["ai_definition"]))
    return "\n".join(parts)


def _examples(item: dict) -> str:
    if "dictionary_senses" not in item or (not item.get("dictionary_senses") and not item.get("example_source") and item.get("example")):
        return str(item.get("example", ""))
    parts = []
    for sense in item.get("dictionary_senses", []):
        for example in sense.get("examples", []):
            parts.append(f"Wordset [{sense.get('pos', '')}]：{example}")
    if not parts:
        parts.append("Wordset 暂无用例")
    if item.get("ai_example"):
        parts.append("AI 生成例句（待核）：" + str(item["ai_example"]))
    return "\n".join(parts)


def _dictionary_license(analysis: dict) -> str:
    if not any("dictionary_senses" in item for item in analysis.get("items", [])):
        return ""
    # Read the fixed bundled notice only; never load a client-supplied path.
    return (Path(__file__).resolve().parents[1] / "data" / "wordset" / "LICENSE").read_text(encoding="utf-8")


def _level_label(item: dict) -> str:
    candidates = sorted({level for level in item.get("level_candidates", []) if level in COLORS})
    if len(candidates) > 1:
        return " / ".join(candidates) + "（义项待核）"
    return str(item.get("level", ""))


def build_docx(analysis: dict) -> bytes:
    document = Document()
    section = document.sections[0]
    section.top_margin = section.bottom_margin = Inches(0.65)
    section.left_margin = section.right_margin = Inches(0.65)
    normal = document.styles["Normal"]
    normal.font.name = "Calibri"
    normal.font.size = Pt(10)
    normal.element.rPr.rFonts.set(qn("w:eastAsia"), "Microsoft YaHei")
    document.add_heading("词境 · LexiScope", 0)
    document.add_paragraph(FORMAT_NOTICE)
    document.add_paragraph("C1：浅黄色高亮 · C2：浅绿色高亮。等级来源与估计标识见词汇表；来源标签不等于当前语境已确认。")
    document.add_heading("标注正文", 1)
    matches = _matches(analysis)
    for segments in _line_segments(analysis["text"], matches):
        paragraph = document.add_paragraph()
        for value, match in segments:
            run = paragraph.add_run(value)
            if match:
                shade = OxmlElement("w:shd")
                shade.set(qn("w:fill"), COLORS[match["item"]["level"]])
                run._r.get_or_add_rPr().append(shade)
                run.bold = True
    document.add_heading("词汇表", 1)
    headers = ["词汇 / 次数", "等级 / 词性", "中文释义", "英文定义", "例句", "来源 / 可靠性"]
    table = document.add_table(rows=1, cols=len(headers))
    table.style = "Table Grid"
    for cell, label in zip(table.rows[0].cells, headers):
        cell.text = label
    repeat_header = OxmlElement("w:tblHeader")
    table.rows[0]._tr.get_or_add_trPr().append(repeat_header)
    for item in analysis.get("items", []):
        values = [
            f'{item.get("lemma", "")}\n出现 {item.get("count", 0)} 次',
            f'{_level_label(item)} / {item.get("pos", "")}',
            item.get("chinese", ""), _definitions(item),
            _examples(item), _basis(item),
        ]
        for cell, value in zip(table.add_row().cells, values):
            cell.text = str(value)
    for warning in analysis.get("warnings", []):
        document.add_paragraph(str(warning))
    if notice := _dictionary_license(analysis):
        document.add_heading("词典资料许可与原始声明", 1)
        document.add_paragraph(notice)
    output = io.BytesIO()
    document.save(output)
    return output.getvalue()


def build_pdf(analysis: dict) -> bytes:
    escape = lambda value: html.escape(str(value), quote=True)
    paragraphs = []
    for segments in _line_segments(analysis["text"], _matches(analysis)):
        runs = []
        for value, match in segments:
            content = escape(value).replace("\t", "&#160;&#160;&#160;&#160;")
            if match:
                content = f'<span style="background-color:#{COLORS[match["item"]["level"]]};font-weight:bold">{content}</span>'
            runs.append(content)
        paragraphs.append("<p>" + ("".join(runs) or "&#160;") + "</p>")
    cards = []
    for item in analysis.get("items", []):
        cards.append(
            '<div class="entry">'
            f'<h3>{escape(item.get("lemma", ""))} · {escape(_level_label(item))} '
            f'· {escape(item.get("pos", ""))} · {escape(item.get("count", 0))} 次</h3>'
            f'<p><b>中文：</b>{escape(item.get("chinese", ""))}</p>'
            f'<p><b>English：</b>{escape(_definitions(item)).replace(chr(10), "<br>")}</p>'
            f'<p><b>例句／短语用例：</b>{escape(_examples(item)).replace(chr(10), "<br>")}</p>'
            f'<p class="source">{escape(_basis(item)).replace(chr(10), "<br>")}</p>'
            '</div>'
        )
    warnings = "".join(f'<p class="source">{escape(warning)}</p>' for warning in analysis.get("warnings", []))
    license_notice = _dictionary_license(analysis)
    license_markup = '<h2>词典资料许可与原始声明</h2><p class="source">' + escape(license_notice).replace(chr(10), "<br>") + '</p>' if license_notice else ''
    markup = (
        '<!DOCTYPE html><html><body><h1>词境 · LexiScope</h1>'
        f'<p class="source">{escape(FORMAT_NOTICE)}</p>'
        '<p>C1：浅黄色 · C2：浅绿色。来源标签不等于当前语境已确认。</p>'
        '<h2>标注正文</h2>' + "".join(paragraphs)
        + '<h2>词汇表</h2>' + "".join(cards) + warnings + license_markup + '</body></html>'
    )
    css = """
    body { font-family: sans-serif; font-size: 10pt; color: #203633; }
    h1 { font-size: 24pt; } h2 { font-size: 16pt; } h3 { font-size: 12pt; }
    p { margin-top: 3pt; margin-bottom: 7pt; }
    .source { font-size: 8pt; color: #52635d; }
    .entry { margin-top: 12pt; border-top: 0.5pt solid #d8dfdb; padding-top: 5pt; }
    """
    story = fitz.Story(html=markup, user_css=css)
    paper = fitz.paper_rect("a4")
    area = fitz.Rect(40, 42, paper.width - 40, paper.height - 42)
    output = io.BytesIO()
    writer = fitz.DocumentWriter(output)
    try:
        more = True
        pages = 0
        while more:
            device = writer.begin_page(paper)
            more, _ = story.place(area)
            story.draw(device)
            writer.end_page()
            pages += 1
            if pages > 500:
                raise ValueError("导出内容过长，请缩短文本或词汇表后重试。")
    finally:
        writer.close()
    return output.getvalue()


def _csv_safe(value) -> str:
    text = str(value if value is not None else "")
    # Excel/LibreOffice may ignore leading whitespace before formula triggers.
    if (text and ord(text[0]) < 32) or re.match(r"^\s*[=+@-]", text):
        return "'" + text
    return text


def build_csv(analysis: dict) -> bytes:
    output = io.StringIO(newline="")
    writer = csv.writer(output)
    license_notice = _dictionary_license(analysis)
    headers = ["词汇", "CEFR", "词性", "中文释义", "英文定义", "例句", "可靠性和来源", "出现次数", "原文词形"]
    writer.writerow(headers + (["词典资料许可与原始声明"] if license_notice else []))
    for item in analysis.get("items", []):
        row = [
            _csv_safe(item.get("lemma")), _csv_safe(_level_label(item)), _csv_safe(item.get("pos")),
            _csv_safe(item.get("chinese")), _csv_safe(_definitions(item)), _csv_safe(_examples(item)),
            _csv_safe(_basis(item)), _csv_safe(item.get("count", 0)), _csv_safe(", ".join(item.get("forms", []))),
        ]
        writer.writerow(row + ([_csv_safe(license_notice)] if license_notice else []))
    return output.getvalue().encode("utf-8-sig")
