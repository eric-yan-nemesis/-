"""Bounded, in-memory DOCX/PDF text extraction. Documents are never retained."""

from __future__ import annotations

import io
import re
import zipfile
from pathlib import PurePosixPath
from urllib.parse import urlsplit
from xml.etree import ElementTree

import fitz
from docx import Document
from docx.table import Table
from docx.text.paragraph import Paragraph

MAX_UPLOAD_BYTES = 10 * 1024 * 1024
MAX_DOCX_EXPANDED_BYTES = 50 * 1024 * 1024
MAX_DOCX_MEMBERS = 2000
MAX_PDF_PAGES = 100
MAX_TEXT_CHARS = 100_000
HYPERLINK_RELATIONSHIPS = {
    "http://schemas.openxmlformats.org/officeDocument/2006/relationships/hyperlink",
    "http://purl.oclc.org/ooxml/officeDocument/relationships/hyperlink",
}


class _InvalidDocument(ValueError):
    pass


def _safe_hyperlink(relationship) -> bool:
    """Allow ordinary link text without ever dereferencing its target."""
    if relationship.attrib.get("Type") not in HYPERLINK_RELATIONSHIPS:
        return False
    try:
        target = urlsplit(relationship.attrib.get("Target", ""))
        return (
            target.scheme.lower() in {"http", "https"}
            and bool(target.hostname)
            and target.username is None
            and target.password is None
        )
    except ValueError:
        return False


def _validate_docx(data: bytes) -> None:
    if not data.startswith(b"PK\x03\x04"):
        raise _InvalidDocument("文件内容不是有效的 DOCX 文档。")
    try:
        with zipfile.ZipFile(io.BytesIO(data)) as archive:
            members = archive.infolist()
            if len(members) > MAX_DOCX_MEMBERS:
                raise _InvalidDocument("DOCX 包含过多文件，无法安全处理。")
            if sum(member.file_size for member in members) > MAX_DOCX_EXPANDED_BYTES:
                raise _InvalidDocument("DOCX 解压后超过 50 MiB 安全限制。")
            names = {member.filename for member in members}
            if not {"[Content_Types].xml", "word/document.xml"}.issubset(names):
                raise _InvalidDocument("DOCX 缺少必需的文档结构。")
            for member in members:
                path = PurePosixPath(member.filename)
                if path.is_absolute() or ".." in path.parts or "\\" in member.filename:
                    raise _InvalidDocument("DOCX 包含不安全的文件路径。")
                if member.flag_bits & 1:
                    raise _InvalidDocument("暂不支持加密的 DOCX 文件。")
                if not member.filename.lower().endswith((".xml", ".rels")):
                    continue
                xml = archive.read(member)
                if re.search(br"<!\s*(?:DOCTYPE|ENTITY)\b", xml, re.IGNORECASE):
                    raise _InvalidDocument("DOCX 包含不允许的 XML 实体或文档类型。")
                root = ElementTree.fromstring(xml)
                if member.filename.lower().endswith(".rels"):
                    for relationship in root:
                        if (
                            relationship.attrib.get("TargetMode", "").lower() == "external"
                            and not _safe_hyperlink(relationship)
                        ):
                            raise _InvalidDocument(
                                "DOCX 含外部图片、对象或不安全链接。请在本地移除这些外部关系后重试。"
                            )
    except _InvalidDocument:
        raise
    except (zipfile.BadZipFile, RuntimeError, OSError, ElementTree.ParseError) as exc:
        raise _InvalidDocument("DOCX 文件已损坏或内部结构无效。") from exc


def _extract_docx(data: bytes) -> dict:
    _validate_docx(data)
    document = Document(io.BytesIO(data))
    blocks = []
    for child in document.element.body:
        kind = child.tag.rsplit("}", 1)[-1]
        if kind == "p":
            blocks.append(Paragraph(child, document).text)
        elif kind == "tbl":
            table = Table(child, document)
            for row in table.rows:
                # Merged cells may appear more than once in python-docx's row API.
                seen_cells = set()
                cells = []
                for cell in row.cells:
                    if cell._tc not in seen_cells:
                        seen_cells.add(cell._tc)
                        cells.append(cell.text)
                blocks.append("\t".join(cells))
    text = "\n".join(blocks)
    warnings = [
        "提取正文和表格文本；页眉、页脚、批注、图片及原始分页不保留。导出为重新排版的标注文档。"
    ]
    if not text.strip():
        warnings.append("未提取到正文文本。DOCX 图片中的文字暂不支持 OCR；可先转换为 PDF。")
    return {"text": text, "warnings": warnings}


def _extract_pdf(data: bytes, ocr: bool) -> dict:
    if not data.lstrip().startswith(b"%PDF-"):
        raise _InvalidDocument("文件内容不是有效的 PDF 文档。")
    texts = []
    empty_pages = []
    ocr_pages = []
    with fitz.open(stream=data, filetype="pdf") as document:
        if document.needs_pass:
            raise _InvalidDocument("暂不支持加密或需要密码的 PDF 文件。")
        if len(document) > MAX_PDF_PAGES:
            raise _InvalidDocument("PDF 超过 100 页限制，请先拆分文件。")
        for index, page in enumerate(document):
            text = page.get_text("text", sort=True).rstrip("\n")
            if not any(character.isalnum() for character in text):
                if ocr:
                    try:
                        text_page = page.get_textpage_ocr(language="eng", dpi=150, full=True)
                        text = page.get_text("text", textpage=text_page, sort=True).rstrip("\n")
                        ocr_pages.append(index + 1)
                    except Exception as exc:
                        raise _InvalidDocument(
                            "PDF OCR 无法运行。请安装 Tesseract 及英文语言数据，或上传含可选中文字的 PDF。"
                        ) from exc
                if not text.strip():
                    empty_pages.append(index + 1)
            texts.append(text)
            if sum(map(len, texts)) + 2 * len(texts) > MAX_TEXT_CHARS:
                raise _InvalidDocument("提取文本超过 100,000 字符限制，请先拆分文件。")
    warnings = [
        "PDF 按阅读顺序提取文本；多栏、连字符和表格可能需手工校正。导出不保留原 PDF 的页面布局。"
    ]
    if ocr_pages:
        warnings.append("已对无文本页执行英文 OCR：" + ", ".join(map(str, ocr_pages)) + "。请核对识别结果。")
    if empty_pages:
        warnings.append(
            "以下页面没有可提取文本：" + ", ".join(map(str, empty_pages))
            + ("。" if ocr else "。可勾选英文 OCR 后重新上传。")
        )
    return {"text": "\n\n".join(texts), "warnings": warnings}


def extract_document(filename: str, data: bytes, ocr: bool = False) -> dict:
    """Return normalized text/warnings or raise a safe, Chinese ValueError."""
    if len(data) > MAX_UPLOAD_BYTES:
        raise ValueError("文件超过 10 MiB 大小限制。")
    if not data:
        raise ValueError("上传文件为空。")
    suffix = PurePosixPath(filename or "").suffix.lower()
    if suffix not in {".docx", ".pdf"}:
        raise ValueError("仅支持 DOCX 和 PDF 文件；旧版 DOC 文件请先转换格式。")
    try:
        result = _extract_docx(data) if suffix == ".docx" else _extract_pdf(data, ocr)
        if len(result["text"]) > MAX_TEXT_CHARS:
            raise _InvalidDocument("提取文本超过 100,000 字符限制，请先拆分文件。")
        return result
    except _InvalidDocument as exc:
        raise ValueError(str(exc)) from exc
    except Exception as exc:
        raise ValueError("无法读取该文档。文件可能已损坏、加密，或不是受支持的 DOCX/PDF 格式。") from exc
