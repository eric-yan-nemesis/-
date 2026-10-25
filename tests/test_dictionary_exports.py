import csv
import io
import unicodedata

import fitz
from docx import Document

from app.exports import build_csv, build_docx, build_pdf


def result(with_examples=True):
    return {
        "text": "Meticulous notes.", "occurrences": [], "warnings": [], "ai_used": True,
        "items": [{
            "id": "v1", "lemma": "meticulous", "level": "C2", "pos": "adjective",
            "chinese": "一丝不苟的", "definition": "Original dictionary definition.",
            "example": "Original example." if with_examples else "", "count": 1, "forms": ["Meticulous"],
            "status": "reference", "source": {"name": "Octanove", "license": "CC BY-SA 4.0"},
            "definition_source": {"name": "Wordset Dictionary", "license": "CC BY-SA 4.0"},
            "dictionary_senses": [{"id": "sense1", "pos": "adjective", "definition": "Original dictionary definition.",
                                   "examples": ["Original example.", "Second phrase"] if with_examples else []}],
            "ai_definition": "An AI explanation.", "ai_example": "An AI generated sentence.",
        }],
    }


def export_texts(analysis):
    document = Document(io.BytesIO(build_docx(analysis)))
    docx_text = "\n".join(p.text for p in document.paragraphs) + "\n" + "\n".join(
        cell.text for table in document.tables for row in table.rows for cell in row.cells)
    with fitz.open(stream=build_pdf(analysis), filetype="pdf") as document:
        pdf_text = "\n".join(page.get_text() for page in document)
    csv_text = build_csv(analysis).decode("utf-8-sig")
    return [unicodedata.normalize("NFKC", value) for value in [docx_text, pdf_text, csv_text]]


def test_dictionary_exports_preserve_all_original_examples_and_separate_ai():
    for text in export_texts(result()):
        assert "Original example." in text
        assert "Second phrase" in text
        assert "Original dictionary definition." in text
        assert "An AI explanation." in text and "AI 辅助定义" in text
        assert "An AI generated sentence." in text and "AI 生成例句" in text
        assert "Wordset Dictionary" in text
        assert "Princeton WordNet 3.0" in text
        assert "CC BY-SA 4.0" in text
        assert "Copyright 2006 by Princeton University" in " ".join(text.split())
        assert "THIS SOFTWARE AND DATABASE IS PROVIDED" in " ".join(text.split())


def test_missing_dictionary_examples_are_disclosed_not_filled_with_ai():
    for text in export_texts(result(False)):
        assert "Wordset 暂无用例" in text
        assert "AI 生成例句" in text
        assert "Original example." not in text


def test_dictionary_csv_keeps_formula_safety_after_combining_senses():
    analysis = result()
    analysis["items"][0]["dictionary_senses"][0]["definition"] = "<img src='https://example.invalid'>"
    analysis["items"][0]["dictionary_senses"][0]["examples"] = ["=HYPERLINK(\"bad\")"]
    rows = list(csv.reader(io.StringIO(build_csv(analysis).decode("utf-8-sig"))))
    assert "=HYPERLINK" in rows[1][5]
    assert not rows[1][5].startswith("=")
