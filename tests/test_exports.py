import csv
import io
import unicodedata

import fitz
from docx import Document
from docx.oxml.ns import qn

from app.exports import build_csv, build_docx, build_pdf


def analysis_result(text="😀 A substantial change.\n\nAmbiguous wording."):
    words = [("substantial", "C1", "重大的"), ("Ambiguous", "C2", "模棱两可的")]
    items, occurrences = [], []
    for index, (word, level, meaning) in enumerate(words):
        if word not in text:
            continue
        item_id = f"v{index}"
        start = text.index(word)
        items.append({
            "id": item_id, "lemma": word.lower(), "level": level, "pos": "adjective",
            "chinese": meaning, "definition": "An original English definition.",
            "example": f"The result was {word.lower()}.",
            "status": "reference" if index == 0 else "estimate",
            "source": {"name": "Test source", "url": "https://example.org/reference"} if index == 0 else None,
            "notes": "Check the intended sense.", "count": 1, "forms": [word],
        })
        occurrences.append({"start": start, "end": start + len(word), "item_id": item_id, "surface": word, "context": text, "uncertain": index == 1})
    return {"text": text, "items": items, "occurrences": occurrences, "warnings": [], "ai_used": False}


def test_docx_preserves_text_paragraphs_unicode_highlights_and_vocabulary():
    analysis = analysis_result()
    data = build_docx(analysis)
    assert data.startswith(b"PK")
    document = Document(io.BytesIO(data))
    paragraphs = [paragraph.text for paragraph in document.paragraphs]
    start = paragraphs.index("标注正文") + 1
    end = paragraphs.index("词汇表")
    assert "\n".join(paragraphs[start:end]) == analysis["text"]
    highlights = [
        (run.text, run._r.rPr.find(qn("w:shd")).get(qn("w:fill")))
        for paragraph in document.paragraphs for run in paragraph.runs
        if run._r.rPr is not None and run._r.rPr.find(qn("w:shd")) is not None
    ]
    assert ("substantial", "FFF0C2") in highlights
    assert ("Ambiguous", "DCECE8") in highlights
    table_text = " ".join(cell.text for row in document.tables[0].rows for cell in row.cells)
    assert "重大的" in table_text and "模棱两可的" in table_text
    assert "估计，未经来源验证" in table_text and "https://example.org/reference" in table_text
    assert "An original English definition." in table_text


def test_pdf_contains_original_english_chinese_annotations_and_basis():
    data = build_pdf(analysis_result())
    assert data.startswith(b"%PDF-")
    document = fitz.open(stream=data, filetype="pdf")
    text = "\n".join(page.get_text() for page in document)
    assert "substantial" in text and "Ambiguous wording." in text
    assert "重大的" in text and "模棱两可的" in text
    assert "估计" in text and "example.org/reference" in text


def test_pdf_paginates_without_dropping_end_of_text():
    text = "\n".join(f"Paragraph {index}: A substantial change." for index in range(160))
    data = build_pdf(analysis_result(text))
    document = fitz.open(stream=data, filetype="pdf")
    assert len(document) > 1
    extracted = "\n".join(page.get_text() for page in document)
    assert "Paragraph 0:" in extracted and "Paragraph 159:" in extracted


def test_exports_ignore_invalid_overlapping_occurrence_ranges():
    analysis = analysis_result()
    analysis["occurrences"] += [
        {"start": -5, "end": 5000, "item_id": "v0"},
        {"start": 0, "end": 1, "item_id": "missing"},
        {"start": 5, "end": 6, "item_id": "v0"},
    ]
    assert build_docx(analysis).startswith(b"PK")
    assert build_pdf(analysis).startswith(b"%PDF-")


def test_pdf_escapes_html_without_loading_external_resources():
    text = '<img src="https://evil.example/tracker"> <script>alert(1)</script>'
    data = build_pdf(analysis_result(text))
    extracted = "".join(page.get_text() for page in fitz.open(stream=data, filetype="pdf"))
    assert "evil.example/tracker" in extracted and "alert(1)" in extracted


def test_csv_bom_roundtrips_chinese_and_quotes():
    analysis = analysis_result()
    analysis["items"][0]["chinese"] = '重要，"尤其重要"\n第二行'
    data = build_csv(analysis)
    assert data.startswith(b"\xef\xbb\xbf")
    rows = list(csv.reader(io.StringIO(data.decode("utf-8-sig"))))
    assert len(rows) == 3
    assert rows[1][3] == '重要，"尤其重要"\n第二行'
    assert rows[2][0] == "ambiguous"


def test_csv_neutralizes_formula_injection_in_every_text_column():
    analysis = analysis_result()
    item = analysis["items"][0]
    item.update({
        "lemma": "=HYPERLINK(\"https://evil.example\")", "level": "+C1", "pos": "-10",
        "chinese": "@SUM(1)", "definition": "  =1+1", "example": "\t=1+1", "forms": ["=CMD()"],
    })
    rows = list(csv.reader(io.StringIO(build_csv(analysis).decode("utf-8-sig"))))
    for column in [0, 1, 2, 3, 4, 5, 8]:
        assert rows[1][column].startswith("'")


def test_all_exports_retain_profile_attribution_license_changes_and_ambiguous_levels():
    analysis = analysis_result()
    item = analysis["items"][0]
    item["source"] = {
        "name": "Octanove Vocabulary Profile C1/C2 v1.0",
        "url": "https://github.com/openlanguageprofiles/olp-en-cefrj",
        "license": "CC BY-SA 4.0",
    }
    item["level_candidates"] = ["C2", "C1"]
    document = Document(io.BytesIO(build_docx(analysis)))
    docx_text = " ".join(cell.text for table in document.tables for row in table.rows for cell in row.cells)
    pdf = fitz.open(stream=build_pdf(analysis), filetype="pdf")
    pdf_text = "\n".join(page.get_text() for page in pdf)
    csv_text = build_csv(analysis).decode("utf-8-sig")
    for text in (docx_text, pdf_text, csv_text):
        # PDF typography may encode common Latin ligatures as single glyphs.
        text = unicodedata.normalize("NFKC", text)
        assert "Octanove Labs / Open Language Profiles" in text
        assert "CC BY-SA 4.0" in text
        assert "https://creativecommons.org/licenses/by-sa/4.0/" in text
        assert "词形匹配" in text
        assert "C1 / C2(义项待核)" in text
