import io
import shutil
import socket
import subprocess
import zipfile

import fitz
import pytest
from docx import Document
from docx.opc.constants import RELATIONSHIP_TYPE
from docx.oxml import OxmlElement
from docx.oxml.ns import qn

from app.documents import extract_document


def docx_bytes(document):
    output = io.BytesIO()
    document.save(output)
    return output.getvalue()


def pdf_bytes(text="A substantial opportunity."):
    document = fitz.open()
    page = document.new_page()
    if text:
        page.insert_text((40, 40), text)
    return document.tobytes()


def test_docx_preserves_paragraph_table_and_following_paragraph_order():
    document = Document()
    document.add_paragraph("First paragraph.\nA line break.")
    table = document.add_table(rows=1, cols=2)
    table.cell(0, 0).text = "Table left"
    table.cell(0, 1).text = "Table right"
    document.add_paragraph("Final paragraph.")
    result = extract_document("notes.DOCX", docx_bytes(document))
    assert result["text"] == "First paragraph.\nA line break.\nTable left\tTable right\nFinal paragraph."
    assert result["warnings"]


def test_pdf_extracts_real_text():
    result = extract_document("notes.pdf", pdf_bytes())
    assert result["text"] == "A substantial opportunity."


@pytest.mark.parametrize("filename,data", [
    ("notes.txt", b"hello"), ("notes.doc", b"hello"),
    ("notes.pdf", b"PK\x03\x04false"), ("notes.docx", b"%PDF-false"),
    ("notes.pdf", b"%PDF-1.7\ncorrupt"), ("notes.docx", b"PK\x03\x04corrupt"),
    ("notes.docx", b""),
])
def test_invalid_types_signatures_and_corrupt_data_fail_safely(filename, data):
    with pytest.raises(ValueError):
        extract_document(filename, data)


def test_raw_file_limit_is_checked_before_parser():
    with pytest.raises(ValueError, match="10 MiB"):
        extract_document("large.pdf", b"x" * (10 * 1024 * 1024 + 1))


def test_docx_zip_bomb_inflated_size_limit():
    output = io.BytesIO()
    with zipfile.ZipFile(output, "w", compression=zipfile.ZIP_DEFLATED) as archive:
        archive.writestr("word/document.xml", b"x" * (50 * 1024 * 1024 + 1))
        archive.writestr("[Content_Types].xml", "<Types/>")
    with pytest.raises(ValueError, match="50 MiB"):
        extract_document("bomb.docx", output.getvalue())


@pytest.mark.parametrize("path,xml", [
    ("word/external.rels", '<Relationships><Relationship TargetMode="External" Target="https://example.com"/></Relationships>'),
    ("word/unsafe.xml", '<!DOCTYPE t [<!ENTITY foo "bar">]><t>&foo;</t>'),
    ("../unsafe.xml", "<t/>"),
])
def test_docx_rejects_external_resources_entities_and_traversal(path, xml):
    document = Document()
    document.add_paragraph("Safe text")
    output = io.BytesIO(docx_bytes(document))
    with zipfile.ZipFile(output, "a") as archive:
        archive.writestr(path, xml)
    with pytest.raises(ValueError):
        extract_document("unsafe.docx", output.getvalue())


def test_docx_allows_safe_hyperlink_text_without_any_network_fetch(monkeypatch):
    document = Document()
    paragraph = document.add_paragraph("Read ")
    relationship = paragraph.part.relate_to(
        "https://example.org/a-reference", RELATIONSHIP_TYPE.HYPERLINK, is_external=True
    )
    hyperlink = OxmlElement("w:hyperlink")
    hyperlink.set(qn("r:id"), relationship)
    run = OxmlElement("w:r")
    text = OxmlElement("w:t")
    text.text = "the dictionary"
    run.append(text)
    hyperlink.append(run)
    paragraph._p.append(hyperlink)

    def no_network(*args, **kwargs):
        raise AssertionError("DOCX extraction must not fetch external hyperlinks")
    monkeypatch.setattr(socket.socket, "connect", no_network)
    result = extract_document("linked.docx", docx_bytes(document))
    assert result["text"] == "Read the dictionary"


@pytest.mark.parametrize("relation_type,target", [
    (RELATIONSHIP_TYPE.IMAGE, "https://example.org/tracker.png"),
    (RELATIONSHIP_TYPE.OLE_OBJECT, "https://example.org/remote-object"),
    (RELATIONSHIP_TYPE.HYPERLINK, "file:///etc/passwd"),
    (RELATIONSHIP_TYPE.HYPERLINK, "javascript:alert(1)"),
    (RELATIONSHIP_TYPE.HYPERLINK, "https://user:password@example.org"),
])
def test_docx_rejects_external_objects_and_unsafe_hyperlinks(relation_type, target):
    document = Document()
    document.add_paragraph("Local text")
    document.part.relate_to(target, relation_type, is_external=True)
    with pytest.raises(ValueError, match="外部"):
        extract_document("unsafe.docx", docx_bytes(document))


def test_encrypted_pdf_rejected():
    document = fitz.open(stream=pdf_bytes(), filetype="pdf")
    encrypted = document.tobytes(encryption=fitz.PDF_ENCRYPT_AES_256, owner_pw="owner", user_pw="reader")
    with pytest.raises(ValueError, match="密码"):
        extract_document("encrypted.pdf", encrypted)


def test_pdf_page_limit():
    document = fitz.open()
    for _ in range(101):
        document.new_page()
    with pytest.raises(ValueError, match="100 页"):
        extract_document("too-many.pdf", document.tobytes())


def test_image_only_pdf_reports_ocr_option_without_silent_success():
    result = extract_document("scan.pdf", pdf_bytes(""))
    assert result["text"] == ""
    assert any("OCR" in warning for warning in result["warnings"])


def test_ocr_is_opt_in_and_only_applied_to_empty_pages(monkeypatch):
    def fail_if_ocr(*args, **kwargs):
        raise AssertionError("OCR must not be used on text-bearing pages")
    monkeypatch.setattr(fitz.Page, "get_textpage_ocr", fail_if_ocr)
    assert extract_document("text.pdf", pdf_bytes(), ocr=True)["text"]
    assert extract_document("scan.pdf", pdf_bytes(""), ocr=False)["text"] == ""


def test_unavailable_ocr_returns_specific_action(monkeypatch):
    def unavailable(*args, **kwargs):
        raise RuntimeError("Missing tessdata")
    monkeypatch.setattr(fitz.Page, "get_textpage_ocr", unavailable)
    with pytest.raises(ValueError, match="Tesseract"):
        extract_document("scan.pdf", pdf_bytes(""), ocr=True)


def test_live_english_ocr_on_image_only_pdf_when_tesseract_is_available():
    executable = shutil.which("tesseract")
    if not executable:
        pytest.skip("Optional Tesseract OCR engine is not installed")
    languages = subprocess.run([executable, "--list-langs"], capture_output=True, text=True, check=True)
    if "eng" not in languages.stdout.splitlines():
        pytest.skip("Optional Tesseract English language data is not installed")
    original = fitz.open()
    page = original.new_page(width=900, height=240)
    page.insert_text((45, 110), "A meticulous researcher is resilient.", fontsize=32)
    image = page.get_pixmap(matrix=fitz.Matrix(2, 2))
    scanned = fitz.open()
    scanned_page = scanned.new_page(width=900, height=240)
    scanned_page.insert_image(scanned_page.rect, stream=image.tobytes("png"))
    data = scanned.tobytes()
    without_ocr = extract_document("scan.pdf", data, ocr=False)
    assert without_ocr["text"] == ""
    assert any("OCR" in warning for warning in without_ocr["warnings"])
    recognized = extract_document("scan.pdf", data, ocr=True)
    assert "meticulous" in recognized["text"].lower()
    assert "resilient" in recognized["text"].lower()
    assert any("已对" in warning for warning in recognized["warnings"])


def test_text_character_limit():
    document = Document()
    document.add_paragraph("x" * 100_001)
    with pytest.raises(ValueError, match="100,000"):
        extract_document("long.docx", docx_bytes(document))
