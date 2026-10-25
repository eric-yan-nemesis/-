import io

import pytest
from docx import Document
from fastapi.testclient import TestClient

from app import ai
from app.main import MAX_BODY, app


@pytest.fixture
def client():
    with TestClient(app) as value:
        yield value


def test_health_and_paste_work_without_ai(client, monkeypatch):
    monkeypatch.delenv('AI_API_KEY', raising=False)
    async def forbidden(*args, **kwargs):
        raise AssertionError('AI must be opt-in')
    monkeypatch.setattr(ai, 'enrich', forbidden)
    response = client.get('/api/health')
    assert response.status_code == 200
    assert response.json()['status'] == 'ok'
    assert response.json()['ai_configured'] is False
    assert response.json()['reference_entries'] > 2000
    response = client.post('/api/analyze', json={'text': 'Meticulous notes alleviate complexity.'})
    assert response.status_code == 200
    assert len(response.json()['occurrences']) == 3
    assert response.headers['Cache-Control'] == 'no-store'
    assert "script-src 'self'" in response.headers['Content-Security-Policy']
    assert response.headers['X-Content-Type-Options'] == 'nosniff'


@pytest.mark.parametrize('text', ['', '  \n ', 'Hello\x00world', 'x' * 100001])
def test_invalid_text_is_rejected(client, text):
    assert client.post('/api/analyze', json={'text': text}).status_code == 422


def test_client_cannot_supply_api_key_or_ai_endpoint(client):
    assert client.post('/api/analyze', json={'text': 'Hello', 'AI_API_KEY': 'not-supported'}).status_code == 422


def test_ai_only_runs_when_explicitly_requested(client, monkeypatch):
    calls = []
    async def enrich(result, include_estimates):
        calls.append(include_estimates)
        return result
    monkeypatch.setattr(ai, 'enrich', enrich)
    response = client.post('/api/analyze', json={'text': 'Meticulous notes.', 'use_ai': True, 'include_estimates': False})
    assert response.status_code == 200 and calls == [False]


def test_cross_site_posts_and_oversized_requests_are_rejected(client):
    assert client.post('/api/analyze', headers={'Origin': 'https://attacker.test'}, json={'text': 'Hello'}).status_code == 403
    response = client.post('/api/analyze', content=b'x' * (MAX_BODY + 1), headers={'Content-Type': 'application/json'})
    assert response.status_code == 413


def test_real_docx_upload_extracts_paragraphs_then_analyzes(client):
    document = Document()
    document.add_paragraph('Meticulous research.')
    document.add_paragraph('A plausible explanation.')
    output = io.BytesIO()
    document.save(output)
    response = client.post('/api/documents', files={'file': ('study.docx', output.getvalue(),
        'application/vnd.openxmlformats-officedocument.wordprocessingml.document')}, data={'ocr': 'false'})
    assert response.status_code == 200
    assert response.json()['text'] == 'Meticulous research.\nA plausible explanation.'
    analyzed = client.post('/api/analyze', json={'text': response.json()['text']})
    assert analyzed.status_code == 200
    assert {'meticulous', 'plausible'} <= {item['lemma'] for item in analyzed.json()['items']}


def test_upload_signature_mismatch_is_rejected(client):
    response = client.post('/api/documents', files={'file': ('fake.pdf', b'not a pdf', 'application/pdf')})
    assert response.status_code == 400


@pytest.mark.parametrize('format', ['docx', 'pdf', 'csv'])
def test_exports_return_real_files(client, format):
    analyzed = client.post('/api/analyze', json={'text': 'Meticulous records.\nA candid report.'}).json()
    response = client.post(f'/api/export/{format}', json={'analysis': analyzed})
    assert response.status_code == 200
    assert f'lexiscope-annotated.{format}' in response.headers['content-disposition']
    assert response.content.startswith({'docx': b'PK', 'pdf': b'%PDF', 'csv': b'\xef\xbb\xbf'}[format])


def test_export_rejects_invalid_highlight_spans(client):
    analyzed = client.post('/api/analyze', json={'text': 'Meticulous notes.'}).json()
    analyzed['occurrences'][0]['end'] = 9000
    response = client.post('/api/export/docx', json={'analysis': analyzed})
    assert response.status_code == 422


def test_export_rejects_unknown_vocabulary_id(client):
    analyzed = client.post('/api/analyze', json={'text': 'Meticulous notes.'}).json()
    analyzed['occurrences'][0]['item_id'] = 'does-not-exist'
    assert client.post('/api/export/csv', json={'analysis': analyzed}).status_code == 422


@pytest.mark.parametrize('content', [
    b'{"text":"private document \\ud800"}',
    b'{"text":"private document \\udfff"}',
    b'{"text":"private document \\uffff"}',
    b'{"text":"private document \\u0000"}',
])
def test_invalid_unicode_is_rejected_without_echoing_document(client, content):
    response = client.post('/api/analyze', content=content, headers={'Content-Type': 'application/json'})
    assert response.status_code == 422
    assert 'private document' not in response.text
    assert response.json()['detail']


def test_export_rejects_xml_invalid_annotation_without_echoing_it(client):
    analyzed = client.post('/api/analyze', json={'text': 'Meticulous notes.'}).json()
    analyzed['items'][0]['chinese'] = 'private annotation\x00'
    response = client.post('/api/export/docx', json={'analysis': analyzed})
    assert response.status_code == 422
    assert 'private annotation' not in response.text
